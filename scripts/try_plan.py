"""
Try the interview planner on a folder of real resumes and write a review sheet.

Place at: scripts/try_plan.py   (run from the project root)
Install:  pip install openai pypdf python-docx      (pypdf / python-docx only for PDF / DOCX files)

  python scripts/try_plan.py resumes/                       # every .pdf .docx .txt .md in the folder
  python scripts/try_plan.py resumes/ --count 16 --out data/plan_review.md

For each resume it shows how much text was read, what your parser found, and the full list of questions the planner
wrote with their key points, so someone from that field can mark them good or bad. Nothing is saved to the app.
Uses the same OPENAI_API_KEY and PLAN_MODEL as the app. Keep resumes private: do not commit the folder or the output.
"""
from __future__ import annotations

import argparse
import importlib
import sys
from pathlib import Path

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

sys.path.insert(0, str(Path.cwd()))


def read_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    if ext == ".docx":
        import docx
        doc = docx.Document(str(path))
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        return "\n".join([p.text for p in doc.paragraphs] + cells)
    raise ValueError(f"unsupported file type {ext}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("folder", type=Path)
    ap.add_argument("--count", type=int, default=16)
    ap.add_argument("--out", type=Path, default=Path("data/plan_review.md"))
    ap.add_argument("--parser-module", default="backend.app.services.resume_parser",
                    help="module that has extract_resume_profile(text)")
    ap.add_argument("--planner-module", default="backend.app.interview.planner")
    args = ap.parse_args()

    parse = importlib.import_module(args.parser_module).extract_resume_profile
    planner = importlib.import_module(args.planner_module)

    files = sorted(p for p in args.folder.iterdir() if p.suffix.lower() in {".pdf", ".docx", ".txt", ".md"})
    if not files:
        sys.exit(f"No resumes found in {args.folder}")
    args.out.parent.mkdir(parents=True, exist_ok=True)

    lines = ["# Interview plan review", "",
             "For each resume mark every question: [x] good, or note what is wrong (too easy, wrong field, vague, unfair).", ""]
    for n, path in enumerate(files, 1):
        print(f"[{n}/{len(files)}] {path.name}")
        try:
            text = read_text(path)
        except Exception as exc:
            lines += [f"## {n}. {path.name}", f"Could not read: {exc}", ""]
            continue
        profile = parse(text)
        resume = {"profile": profile, "text": text}
        plan = planner.build_plan(resume, None, args.count)
        lines += [f"## {n}. {path.name}", "",
                  f"- Text read: {len(text)} characters. Sections found by the parser: {', '.join(profile.get('sections_found') or []) or 'none'}",
                  f"- Parser skills: {', '.join(profile.get('skills', [])[:15]) or 'none'}",
                  f"- Questions written: {len(plan)}", ""]
        if not plan:
            lines += ["**No questions were written** (too little text, or the model call failed).", ""]
        for i, q in enumerate(plan, 1):
            lines += [f"{i}. [ ] **{q['question']}**",
                      f"   - {q['plan_kind']} | {q['difficulty']} | {q['topic']} / {q['subtopic']}",
                      f"   - Strong answer: {'; '.join(q['key_points'])}", ""]
    args.out.write_text("\n".join(lines), encoding="utf-8")
    print(f"Done. Open {args.out}")


if __name__ == "__main__":
    main()